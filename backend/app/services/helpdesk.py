"""Helpdesk queries. Access decisions use permission names stored in the database."""
from __future__ import annotations

import json
import re
import uuid
from datetime import datetime, timedelta
from pathlib import Path
from zoneinfo import ZoneInfo

from fastapi import UploadFile
from sqlalchemy.orm import Session, joinedload

from app.core.config import settings
from app.core.deps_permission import permission_name_set
from app.core.exceptions import BadRequestException, ForbiddenException, NotFoundException
from app.models.auth import Permission, Role, User, role_has_permissions, user_roles
from app.models.common import Setting
from app.models.helpdesk import (
    Article,
    Category,
    ManualPage,
    Notification,
    Service,
    Ticket,
    TicketAttachment,
    TicketReply,
)

JAKARTA = ZoneInfo("Asia/Jakarta")
DEFAULT_MIMES = [
    "application/pdf",
    "image/png",
    "image/jpeg",
    "image/webp",
    "application/msword",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
]


def setting_value(db: Session, name: str, default):
    row = db.query(Setting).filter(Setting.name == name).first()
    if row is None or row.payload is None:
        return default
    return row.payload


def helpdesk_config(db: Session) -> dict:
    max_chars = int(setting_value(db, "ticket_description_max", 150))
    max_bytes = int(setting_value(db, "ticket_attachment_max_bytes", settings.MAX_UPLOAD_SIZE))
    mimes = setting_value(db, "ticket_attachment_mimes", DEFAULT_MIMES)
    return {
        "description_max": max_chars,
        "attachment_max_bytes": max_bytes,
        "attachment_mimes": mimes,
    }


def _names(db: Session, user: User) -> set[str]:
    return permission_name_set(db, user)


def sees_all_tickets(names: set[str]) -> bool:
    return "view_all_ticket" in names


def can_view_ticket(names: set[str], user: User, ticket: Ticket) -> bool:
    if "view_all_ticket" in names:
        return True
    return "view_own_ticket" in names and ticket.requester_id == user.id


def can_reply(names: set[str], user: User, ticket: Ticket) -> bool:
    if ticket.status == "resolved":
        return False
    if "reply_ticket" in names:
        return True
    return "reply_own_ticket" in names and ticket.requester_id == user.id


def _ticket_query(db: Session):
    return db.query(Ticket).options(
        joinedload(Ticket.service),
        joinedload(Ticket.category),
        joinedload(Ticket.requester),
    )


def scoped_tickets(db: Session, user: User):
    names = _names(db, user)
    query = _ticket_query(db)
    if "view_all_ticket" in names:
        return query
    if "view_own_ticket" in names:
        return query.filter(Ticket.requester_id == user.id)
    return query.filter(Ticket.id == -1)


def summary(db: Session, user: User) -> dict:
    rows = scoped_tickets(db, user).all()
    latest = sorted(rows, key=lambda t: t.created_at or datetime.min, reverse=True)[:8]
    by_institution: dict[str, dict] = {}
    for ticket in rows:
        name = (ticket.requester.institution if ticket.requester and ticket.requester.institution else None) or "—"
        bucket = by_institution.setdefault(
            name,
            {"name": name, "total": 0, "open": 0, "in_progress": 0, "resolved": 0},
        )
        bucket["total"] += 1
        if ticket.status in ("open", "in_progress", "resolved"):
            bucket[ticket.status] += 1
    institutions = sorted(by_institution.values(), key=lambda item: (-item["total"], item["name"]))
    return {
        "total": len(rows),
        "in_progress": sum(1 for t in rows if t.status == "in_progress"),
        "resolved": sum(1 for t in rows if t.status == "resolved"),
        "open": sum(1 for t in rows if t.status == "open"),
        "institution_count": len([item for item in institutions if item["name"] != "—"]),
        "by_institution": institutions,
        "recent": [ticket_list_item(t) for t in latest],
        "scope": "all" if sees_all_tickets(_names(db, user)) else "own",
    }


def ticket_list_item(ticket: Ticket) -> dict:
    requester = ticket.requester
    return {
        "id": ticket.uuid,
        "title": ticket.title,
        "status": ticket.status,
        "service_name": ticket.service.name if ticket.service else "",
        "category_name": ticket.category.name if ticket.category else "",
        "staff_priority": bool(ticket.staff_priority),
        "requester_name": requester.fullname if requester else None,
        "institution": requester.institution if requester else None,
        "study_program": ticket.study_program,
        "created_at": ticket.created_at,
        "updated_at": ticket.updated_at,
    }


def attachment_item(row: TicketAttachment) -> dict:
    return {
        "id": row.id,
        "file_name": row.file_name,
        "mime_type": row.mime_type,
        "size": row.size,
        "created_at": row.created_at,
    }


def ticket_detail(db: Session, ticket: Ticket) -> dict:
    replies = (
        db.query(TicketReply)
        .options(joinedload(TicketReply.author), joinedload(TicketReply.attachments))
        .filter(TicketReply.ticket_id == ticket.id)
        .order_by(TicketReply.created_at.asc(), TicketReply.id.asc())
        .all()
    )
    own_files = (
        db.query(TicketAttachment)
        .filter(TicketAttachment.ticket_id == ticket.id, TicketAttachment.reply_id.is_(None))
        .all()
    )
    data = ticket_list_item(ticket)
    data.update(
        {
            "description": ticket.description,
            "service_id": ticket.service_id,
            "category_id": ticket.category_id,
            "requester_id": ticket.requester_id,
            "resolved_at": ticket.resolved_at,
            "attachments": [attachment_item(a) for a in own_files],
            "replies": [
                {
                    "id": reply.id,
                    "body": reply.body,
                    "author_id": reply.author_id,
                    "author_name": reply.author.fullname if reply.author else "",
                    "created_at": reply.created_at,
                    "attachments": [attachment_item(a) for a in (reply.attachments or [])],
                }
                for reply in replies
            ],
        }
    )
    return data


def get_ticket_or_404(db: Session, user: User, ticket_ref: str | int) -> Ticket:
    """Resolve ticket by public UUID (preferred) or legacy numeric id."""
    query = _ticket_query(db)
    ref = str(ticket_ref).strip()
    if ref.isdigit():
        ticket = query.filter(Ticket.id == int(ref)).first()
    else:
        ticket = query.filter(Ticket.uuid == ref).first()
    if not ticket or not can_view_ticket(_names(db, user), user, ticket):
        raise NotFoundException("Tiket tidak ditemukan")
    return ticket


def _private_dir() -> Path:
    base = Path(settings.UPLOAD_DIR)
    if not base.is_absolute():
        base = Path(__file__).resolve().parents[2] / base
    folder = base / "private" / "tickets"
    folder.mkdir(parents=True, exist_ok=True)
    return folder


async def store_uploads(
    db: Session,
    files: list[UploadFile],
    *,
    ticket_id: int | None,
    reply_id: int | None,
) -> None:
    if not files:
        return
    config = helpdesk_config(db)
    allowed = set(config["attachment_mimes"])
    max_bytes = int(config["attachment_max_bytes"])
    folder = _private_dir()
    for upload in files:
        if not upload or not upload.filename:
            continue
        content = await upload.read()
        mime = upload.content_type or "application/octet-stream"
        if mime not in allowed:
            raise BadRequestException(f"Tipe file tidak diizinkan: {upload.filename}")
        if len(content) > max_bytes:
            raise BadRequestException(f"Ukuran file melebihi batas: {upload.filename}")
        stored = f"{uuid.uuid4().hex}{Path(upload.filename).suffix.lower()}"
        (folder / stored).write_bytes(content)
        db.add(
            TicketAttachment(
                ticket_id=ticket_id,
                reply_id=reply_id,
                file_name=Path(upload.filename).name[:255],
                stored_name=stored,
                mime_type=mime,
                size=len(content),
                created_at=datetime.utcnow(),
            )
        )


def attachment_path(row: TicketAttachment) -> Path:
    return _private_dir() / row.stored_name


def validate_description(db: Session, description: str) -> str:
    text = (description or "").strip()
    limit = int(helpdesk_config(db)["description_max"])
    if not text:
        raise BadRequestException("Deskripsi wajib diisi")
    if len(text) > limit:
        raise BadRequestException(f"Deskripsi maksimal {limit} karakter")
    return text


def create_ticket(
    db: Session,
    user: User,
    *,
    title: str,
    service_id: int,
    category_id: int,
    staff_priority: bool,
    description: str,
    study_program: str | None = None,
) -> Ticket:
    names = _names(db, user)
    if "create_ticket" not in names:
        raise ForbiddenException("Insufficient permissions")
    clean_title = (title or "").strip()
    if not clean_title:
        raise BadRequestException("Judul wajib diisi")
    clean_program = (study_program or "").strip()
    if not clean_program:
        raise BadRequestException("Jurusan wajib diisi")
    service = db.query(Service).filter(Service.id == service_id, Service.is_active.is_(True)).first()
    if not service:
        raise BadRequestException("Layanan tidak ditemukan")
    category = (
        db.query(Category)
        .filter(Category.id == category_id, Category.service_id == service_id)
        .first()
    )
    if not category:
        raise BadRequestException("Kategori masalah tidak sesuai dengan layanan")
    ticket = Ticket(
        uuid=str(uuid.uuid4()),
        title=clean_title[:255],
        description=validate_description(db, description),
        study_program=clean_program[:255],
        service_id=service_id,
        category_id=category_id,
        staff_priority=staff_priority,
        status="open",
        requester_id=user.id,
    )
    db.add(ticket)
    db.flush()
    return ticket


def add_reply(db: Session, user: User, ticket: Ticket, body: str) -> TicketReply:
    names = _names(db, user)
    if not can_reply(names, user, ticket):
        raise ForbiddenException("Insufficient permissions")
    text = (body or "").strip()
    if not text:
        raise BadRequestException("Balasan wajib diisi")
    reply = TicketReply(ticket_id=ticket.id, author_id=user.id, body=text)
    db.add(reply)
    db.flush()
    if "reply_ticket" in names and ticket.requester_id != user.id:
        if ticket.status == "open":
            ticket.status = "in_progress"
        if not ticket.assigned_to:
            ticket.assigned_to = user.id
        notify(
            db,
            ticket.requester_id,
            "ticket.replied",
            {"ticket_id": ticket.uuid, "title": ticket.title},
        )
    else:
        notify_permission(
            db,
            "reply_ticket",
            "ticket.requester_replied",
            {"ticket_id": ticket.uuid, "title": ticket.title},
            exclude_user_id=user.id,
        )
    ticket.updated_at = datetime.utcnow()
    return reply


def close_ticket(db: Session, user: User, ticket: Ticket) -> Ticket:
    names = _names(db, user)
    if "close_ticket" not in names:
        raise ForbiddenException("Insufficient permissions")
    if ticket.status == "resolved":
        raise BadRequestException("Tiket sudah diselesaikan")
    ticket.status = "resolved"
    ticket.resolved_at = datetime.utcnow()
    ticket.updated_at = datetime.utcnow()
    notify(
        db,
        ticket.requester_id,
        "ticket.status",
        {"ticket_id": ticket.uuid, "title": ticket.title, "status": "resolved"},
    )
    return ticket


def delete_ticket(db: Session, user: User, ticket: Ticket) -> str:
    names = _names(db, user)
    if "delete_ticket" not in names:
        raise ForbiddenException("Insufficient permissions")
    ticket_uuid = ticket.uuid
    db.delete(ticket)
    return ticket_uuid


def notify(db: Session, user_id: str, message_key: str, payload: dict) -> None:
    db.add(
        Notification(
            user_id=user_id,
            message_key=message_key,
            payload=json.dumps(payload),
            is_read=False,
            created_at=datetime.utcnow(),
        )
    )


def notify_permission(
    db: Session,
    permission_name: str,
    message_key: str,
    payload: dict,
    exclude_user_id: str | None = None,
) -> None:
    users = (
        db.query(User)
        .join(user_roles, user_roles.c.user_id == User.id)
        .join(Role, Role.id == user_roles.c.role_id)
        .join(role_has_permissions, role_has_permissions.c.role_id == Role.id)
        .join(Permission, Permission.id == role_has_permissions.c.permission_id)
        .filter(Permission.name == permission_name, User.deleted_at.is_(None))
        .distinct()
        .all()
    )
    for person in users:
        if exclude_user_id and person.id == exclude_user_id:
            continue
        notify(db, person.id, message_key, payload)


def list_articles(db: Session, q: str | None, published_only: bool) -> list[Article]:
    query = db.query(Article)
    if published_only:
        query = query.filter(Article.status == "published")
    if q:
        like = f"%{q.strip()}%"
        query = query.filter((Article.title.ilike(like)) | (Article.body.ilike(like)))
    return query.order_by(Article.created_at.desc()).all()


def _article_pdf_dir() -> Path:
    base = Path(settings.UPLOAD_DIR)
    folder = base / "private" / "articles"
    folder.mkdir(parents=True, exist_ok=True)
    return folder


def article_pdf_path(row: Article) -> Path | None:
    if not row.pdf_stored_name:
        return None
    return _article_pdf_dir() / row.pdf_stored_name


def _clear_article_pdf(row: Article) -> None:
    path = article_pdf_path(row)
    if path and path.exists():
        path.unlink(missing_ok=True)
    row.pdf_file_name = None
    row.pdf_stored_name = None
    row.pdf_mime_type = None
    row.pdf_size = None


async def store_article_pdf(db: Session, row: Article, upload: UploadFile | None) -> None:
    if not upload or not upload.filename:
        return
    content = await upload.read()
    mime = upload.content_type or "application/octet-stream"
    if mime != "application/pdf" and not upload.filename.lower().endswith(".pdf"):
        raise BadRequestException("Hanya file PDF yang diizinkan")
    max_bytes = int(helpdesk_config(db)["attachment_max_bytes"])
    if len(content) > max_bytes:
        raise BadRequestException("Ukuran PDF melebihi batas")
    _clear_article_pdf(row)
    stored = f"{uuid.uuid4().hex}.pdf"
    (_article_pdf_dir() / stored).write_bytes(content)
    row.pdf_file_name = Path(upload.filename).name[:255]
    row.pdf_stored_name = stored
    row.pdf_mime_type = "application/pdf"
    row.pdf_size = len(content)


def article_item(row: Article) -> dict:
    return {
        "id": row.id,
        "title": row.title,
        "body": row.body,
        "status": row.status,
        "card_color": row.card_color,
        "created_at": row.created_at,
        "has_pdf": bool(row.pdf_stored_name),
        "pdf_file_name": row.pdf_file_name,
        "pdf_size": row.pdf_size,
    }


def create_article(
    db: Session,
    *,
    title: str,
    body: str,
    status: str = "published",
    card_color: str = "#2B7DE0",
    has_pdf: bool = False,
) -> Article:
    title = (title or "").strip()
    body = (body or "").strip()
    if not title:
        raise BadRequestException("Judul wajib diisi")
    if not body and not has_pdf:
        raise BadRequestException("Isi artikel atau file PDF wajib diisi")
    if status not in {"draft", "published"}:
        raise BadRequestException("Status tidak valid")
    color = (card_color or "#2B7DE0").strip()
    if not color.startswith("#") or len(color) not in {4, 7}:
        color = "#2B7DE0"
    row = Article(title=title, body=body or "", status=status, card_color=color)
    db.add(row)
    db.commit()
    db.refresh(row)
    return row


def update_article(
    db: Session,
    article_id: int,
    *,
    title: str,
    body: str,
    status: str,
    card_color: str,
    remove_pdf: bool = False,
    has_new_pdf: bool = False,
) -> Article:
    row = db.query(Article).filter(Article.id == article_id).first()
    if not row:
        raise NotFoundException("Artikel tidak ditemukan")
    title = (title or "").strip()
    body = (body or "").strip()
    if not title:
        raise BadRequestException("Judul wajib diisi")
    if remove_pdf and not has_new_pdf:
        _clear_article_pdf(row)
    will_have_pdf = bool(row.pdf_stored_name) or has_new_pdf
    if not body and not will_have_pdf:
        raise BadRequestException("Isi artikel atau file PDF wajib diisi")
    if status not in {"draft", "published"}:
        raise BadRequestException("Status tidak valid")
    color = (card_color or row.card_color or "#2B7DE0").strip()
    if not color.startswith("#") or len(color) not in {4, 7}:
        color = row.card_color or "#2B7DE0"
    row.title = title
    row.body = body or ""
    row.status = status
    row.card_color = color
    db.commit()
    db.refresh(row)
    return row


def delete_article(db: Session, article_id: int) -> None:
    row = db.query(Article).filter(Article.id == article_id).first()
    if not row:
        raise NotFoundException("Artikel tidak ditemukan")
    _clear_article_pdf(row)
    db.delete(row)
    db.commit()


def _slugify(value: str) -> str:
    text = (value or "").strip().lower()
    text = re.sub(r"[^a-z0-9\s-]", "", text)
    text = re.sub(r"[\s_-]+", "-", text).strip("-")
    return text or f"manual-{uuid.uuid4().hex[:8]}"


def _manual_pdf_dir() -> Path:
    base = Path(settings.UPLOAD_DIR)
    folder = base / "private" / "manuals"
    folder.mkdir(parents=True, exist_ok=True)
    return folder


def manual_pdf_path(row: ManualPage) -> Path | None:
    if not row.pdf_stored_name:
        return None
    return _manual_pdf_dir() / row.pdf_stored_name


def _clear_manual_pdf(row: ManualPage) -> None:
    path = manual_pdf_path(row)
    if path and path.exists():
        path.unlink(missing_ok=True)
    row.pdf_file_name = None
    row.pdf_stored_name = None
    row.pdf_mime_type = None
    row.pdf_size = None


async def store_manual_pdf(db: Session, row: ManualPage, upload: UploadFile | None) -> None:
    if not upload or not upload.filename:
        return
    content = await upload.read()
    mime = upload.content_type or "application/octet-stream"
    if mime != "application/pdf" and not upload.filename.lower().endswith(".pdf"):
        raise BadRequestException("Hanya file PDF yang diizinkan")
    max_bytes = int(helpdesk_config(db)["attachment_max_bytes"])
    if len(content) > max_bytes:
        raise BadRequestException("Ukuran PDF melebihi batas")
    _clear_manual_pdf(row)
    stored = f"{uuid.uuid4().hex}.pdf"
    (_manual_pdf_dir() / stored).write_bytes(content)
    row.pdf_file_name = Path(upload.filename).name[:255]
    row.pdf_stored_name = stored
    row.pdf_mime_type = "application/pdf"
    row.pdf_size = len(content)


def manual_item(row: ManualPage, *, include_body: bool = True) -> dict:
    data = {
        "id": row.id,
        "title": row.title,
        "slug": row.slug,
        "sort_order": row.sort_order,
        "has_pdf": bool(row.pdf_stored_name),
        "pdf_file_name": row.pdf_file_name,
        "pdf_size": row.pdf_size,
    }
    if include_body:
        data["body"] = row.body
    return data


def list_manual_pages(db: Session) -> list[ManualPage]:
    return db.query(ManualPage).order_by(ManualPage.sort_order.asc(), ManualPage.id.asc()).all()


def create_manual(
    db: Session,
    *,
    title: str,
    body: str,
    slug: str | None = None,
    sort_order: int = 0,
    has_pdf: bool = False,
) -> ManualPage:
    title = (title or "").strip()
    body = (body or "").strip()
    if not title:
        raise BadRequestException("Judul wajib diisi")
    if not body and not has_pdf:
        raise BadRequestException("Isi manual atau file PDF wajib diisi")
    base = _slugify(slug or title)
    candidate = base
    n = 2
    while db.query(ManualPage).filter(ManualPage.slug == candidate).first():
        candidate = f"{base}-{n}"
        n += 1
    row = ManualPage(title=title, slug=candidate, body=body or "", sort_order=int(sort_order or 0))
    db.add(row)
    db.commit()
    db.refresh(row)
    return row


def update_manual(
    db: Session,
    manual_id: int,
    *,
    title: str,
    body: str,
    slug: str | None = None,
    sort_order: int = 0,
    remove_pdf: bool = False,
    has_new_pdf: bool = False,
) -> ManualPage:
    row = db.query(ManualPage).filter(ManualPage.id == manual_id).first()
    if not row:
        raise NotFoundException("Halaman manual tidak ditemukan")
    title = (title or "").strip()
    body = (body or "").strip()
    if not title:
        raise BadRequestException("Judul wajib diisi")
    if remove_pdf and not has_new_pdf:
        _clear_manual_pdf(row)
    will_have_pdf = bool(row.pdf_stored_name) or has_new_pdf
    if not body and not will_have_pdf:
        raise BadRequestException("Isi manual atau file PDF wajib diisi")
    base = _slugify(slug or title)
    candidate = base
    n = 2
    while True:
        exists = db.query(ManualPage).filter(ManualPage.slug == candidate, ManualPage.id != manual_id).first()
        if not exists:
            break
        candidate = f"{base}-{n}"
        n += 1
    row.title = title
    row.body = body or ""
    row.slug = candidate
    row.sort_order = int(sort_order or 0)
    db.commit()
    db.refresh(row)
    return row


def delete_manual(db: Session, manual_id: int) -> None:
    row = db.query(ManualPage).filter(ManualPage.id == manual_id).first()
    if not row:
        raise NotFoundException("Halaman manual tidak ditemukan")
    _clear_manual_pdf(row)
    db.delete(row)
    db.commit()


_STATUS_LABEL = {
    "open": "Baru",
    "in_progress": "Dalam Proses",
    "resolved": "Diselesaikan",
}


def _aware_jakarta(value: datetime | None) -> datetime | None:
    if value is None:
        return None
    if value.tzinfo is None:
        return value.replace(tzinfo=ZoneInfo("UTC")).astimezone(JAKARTA)
    return value.astimezone(JAKARTA)


def _university_label(institution: str | None) -> str:
    if not institution:
        return ""
    parts = re.split(r"\s*[—–\-−]\s*", institution.strip(), maxsplit=1)
    return (parts[0] or "").strip()


def _format_dt(value: datetime | None) -> str:
    aware = _aware_jakarta(value)
    if not aware:
        return ""
    return aware.strftime("%d/%m/%Y %H:%M")


def _report_tickets_query(db: Session, user: User):
    names = _names(db, user)
    opts = (
        joinedload(Ticket.category),
        joinedload(Ticket.service),
        joinedload(Ticket.requester),
    )
    if "view_all_report" in names:
        return db.query(Ticket).options(*opts), "all"
    if "view_own_report" in names:
        return (
            db.query(Ticket).options(*opts).filter(Ticket.requester_id == user.id),
            "own",
        )
    raise ForbiddenException("Insufficient permissions")


def report_export_rows(db: Session, user: User, days: int) -> list[dict]:
    """Flat ticket rows for PDF/Excel export (created within period)."""
    query, _scope = _report_tickets_query(db, user)
    now = datetime.now(JAKARTA)
    start = (now - timedelta(days=days - 1)).replace(hour=0, minute=0, second=0, microsecond=0)
    tickets = query.order_by(Ticket.created_at.desc()).all()
    rows: list[dict] = []
    for ticket in tickets:
        created = _aware_jakarta(ticket.created_at)
        if not created or created < start:
            continue
        requester = ticket.requester
        rows.append(
            {
                "id": ticket.uuid,
                "title": ticket.title or "",
                "university": _university_label(requester.institution if requester else None),
                "study_program": (ticket.study_program or "").strip(),
                "service": ticket.service.name if ticket.service else "",
                "category": ticket.category.name if ticket.category else "",
                "status": _STATUS_LABEL.get(ticket.status or "", ticket.status or ""),
                "requester": requester.fullname if requester else "",
                "created_at": _format_dt(ticket.created_at),
                "resolved_at": _format_dt(ticket.resolved_at),
            }
        )
    return rows


def report(db: Session, user: User, days: int) -> dict:
    query, scope = _report_tickets_query(db, user)
    now = datetime.now(JAKARTA)
    start = (now - timedelta(days=days - 1)).replace(hour=0, minute=0, second=0, microsecond=0)
    tickets = query.all()

    weeks = []
    cursor = start - timedelta(days=start.weekday())
    end = now
    while cursor <= end:
        week_end = cursor + timedelta(days=6)
        label = f"{cursor.strftime('%b%d')}–{(week_end if week_end < end else end).strftime('%b%d')}"
        created_count = 0
        resolved_count = 0
        for ticket in tickets:
            created = _aware_jakarta(ticket.created_at)
            resolved = _aware_jakarta(ticket.resolved_at)
            if created and cursor <= created <= week_end.replace(hour=23, minute=59, second=59):
                if created >= start:
                    created_count += 1
            if resolved and cursor <= resolved <= week_end.replace(hour=23, minute=59, second=59):
                if resolved >= start:
                    resolved_count += 1
        weeks.append({"label": label, "created": created_count, "resolved": resolved_count})
        cursor = cursor + timedelta(days=7)

    categories: dict[str, int] = {}
    total = 0
    for ticket in tickets:
        created = _aware_jakarta(ticket.created_at)
        if created and created >= start:
            total += 1
            name = ticket.category.name if ticket.category else "Lainnya"
            categories[name] = categories.get(name, 0) + 1

    return {
        "range_days": days,
        "updated_at": datetime.utcnow(),
        "total": total,
        "weeks": weeks,
        "categories": [{"name": name, "count": count} for name, count in categories.items()],
        "scope": scope,
        "rows": report_export_rows(db, user, days),
    }


def _pdf_font_paths() -> tuple[str | None, str | None]:
    candidates = [
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
        Path("C:/Windows/Fonts/arial.ttf"),
        Path("C:/Windows/Fonts/arialbd.ttf"),
    ]
    regular = next((str(p) for p in candidates if p.exists() and "Bold" not in p.name and "bd" not in p.name.lower()), None)
    bold = next((str(p) for p in candidates if p.exists() and ("Bold" in p.name or "bd" in p.name.lower())), None)
    if not regular:
        regular = next((str(p) for p in candidates if p.exists()), None)
    return regular, bold or regular


def build_report_pdf(data: dict) -> bytes:
    from fpdf import FPDF

    rows = data.get("rows") or []
    pdf = FPDF(orientation="L", format="A4")
    pdf.set_auto_page_break(auto=True, margin=12)
    pdf.add_page()
    pdf.set_margins(10, 10, 10)

    regular, bold = _pdf_font_paths()
    if regular:
        pdf.add_font("Report", "", regular)
        pdf.add_font("Report", "B", bold or regular)
        font = "Report"
    else:
        font = "Helvetica"

    pdf.set_font(font, "B", 16)
    pdf.cell(0, 9, "Laporan Tiket — Helpdesk LAMDEPILAR", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font(font, "", 10)
    pdf.cell(0, 6, f"Periode: {data['range_days']} hari terakhir", new_x="LMARGIN", new_y="NEXT")
    pdf.cell(0, 6, f"Total tiket dibuat: {data['total']}  |  Baris detail: {len(rows)}", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(2)

    # Ringkasan kategori
    pdf.set_font(font, "B", 11)
    pdf.cell(0, 7, "Ringkasan kategori", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font(font, "", 9)
    if not data.get("categories"):
        pdf.cell(0, 6, "Tidak ada data kategori pada periode ini.", new_x="LMARGIN", new_y="NEXT")
    else:
        for row in data["categories"]:
            pdf.cell(0, 5.5, f"• {row['name']}: {row['count']}", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(3)

    # Tabel detail
    pdf.set_font(font, "B", 11)
    pdf.cell(0, 7, "Rincian tiket", new_x="LMARGIN", new_y="NEXT")

    headers = ["No", "Judul", "Universitas", "Jurusan", "Layanan", "Kategori", "Status", "Pemohon", "Tanggal", "Selesai"]
    widths = [10, 48, 32, 32, 24, 28, 22, 28, 26, 26]
    line_h = 5.2

    def draw_header():
        pdf.set_font(font, "B", 8)
        pdf.set_fill_color(139, 21, 56)
        pdf.set_text_color(255, 255, 255)
        for title, width in zip(headers, widths):
            pdf.cell(width, 7, title, border=1, fill=True, align="C")
        pdf.ln()
        pdf.set_text_color(30, 30, 30)
        pdf.set_font(font, "", 7.5)

    draw_header()

    if not rows:
        pdf.set_font(font, "", 9)
        pdf.cell(0, 8, "Tidak ada tiket pada periode ini.", new_x="LMARGIN", new_y="NEXT")
    else:
        fill = False
        for idx, row in enumerate(rows, start=1):
            values = [
                str(idx),
                row["title"],
                row["university"] or "-",
                row["study_program"] or "-",
                row["service"] or "-",
                row["category"] or "-",
                row["status"] or "-",
                row["requester"] or "-",
                row["created_at"] or "-",
                row["resolved_at"] or "-",
            ]
            title_w = max(widths[1] - 2, 1)
            title_lines = max(1, int(pdf.get_string_width(values[1]) / title_w) + 1)
            row_h = max(line_h + 1, min(14, line_h * title_lines + 1))
            if pdf.get_y() + row_h > pdf.h - 14:
                pdf.add_page()
                draw_header()
            if fill:
                pdf.set_fill_color(248, 242, 244)
            else:
                pdf.set_fill_color(255, 255, 255)
            x_start = pdf.l_margin
            y0 = pdf.get_y()
            x = x_start
            for value, width in zip(values, widths):
                style = "DF" if fill else "D"
                pdf.rect(x, y0, width, row_h, style=style)
                text = value if len(value) <= 90 else value[:87] + "..."
                pdf.set_xy(x + 0.7, y0 + 0.8)
                pdf.multi_cell(width - 1.4, line_h, text, border=0)
                x += width
            pdf.set_xy(x_start, y0 + row_h)
            fill = not fill

    return bytes(pdf.output())


def build_report_xlsx(data: dict) -> bytes:
    from io import BytesIO

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    rows = data.get("rows") or []
    wb = Workbook()

    # Sheet 1: Rincian
    ws = wb.active
    ws.title = "Rincian Tiket"
    headers = [
        "No",
        "ID",
        "Judul",
        "Universitas",
        "Jurusan",
        "Layanan",
        "Kategori",
        "Status",
        "Pemohon",
        "Tanggal Dibuat",
        "Tanggal Selesai",
    ]
    header_fill = PatternFill("solid", fgColor="8B1538")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    thin = Border(
        left=Side(style="thin", color="D0D5DD"),
        right=Side(style="thin", color="D0D5DD"),
        top=Side(style="thin", color="D0D5DD"),
        bottom=Side(style="thin", color="D0D5DD"),
    )
    alt_fill = PatternFill("solid", fgColor="F8F2F4")

    ws.append([f"Laporan Tiket — Helpdesk LAMDEPILAR"])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    ws["A1"].font = Font(bold=True, size=14, color="8B1538")
    ws.append([f"Periode: {data['range_days']} hari terakhir | Total dibuat: {data['total']}"])
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
    ws["A2"].font = Font(size=10, color="667085")
    ws.append([])

    header_row = 4
    for col, title in enumerate(headers, start=1):
        cell = ws.cell(header_row, col, title)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin

    for idx, row in enumerate(rows, start=1):
        values = [
            idx,
            row["id"],
            row["title"],
            row["university"] or "—",
            row["study_program"] or "—",
            row["service"] or "—",
            row["category"] or "—",
            row["status"] or "—",
            row["requester"] or "—",
            row["created_at"] or "—",
            row["resolved_at"] or "—",
        ]
        r = header_row + idx
        for col, value in enumerate(values, start=1):
            cell = ws.cell(r, col, value)
            cell.border = thin
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if idx % 2 == 0:
                cell.fill = alt_fill

    widths = [6, 8, 36, 28, 28, 18, 22, 14, 20, 16, 16]
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[header_row].height = 22
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(headers))}{header_row + max(len(rows), 1)}"

    # Sheet 2: Ringkasan kategori
    ws2 = wb.create_sheet("Ringkasan Kategori")
    ws2.append(["Kategori", "Jumlah"])
    for col in (1, 2):
        cell = ws2.cell(1, col)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = thin
    for i, row in enumerate(data.get("categories") or [], start=2):
        ws2.cell(i, 1, row["name"]).border = thin
        ws2.cell(i, 2, row["count"]).border = thin
    if not data.get("categories"):
        ws2.cell(2, 1, "Tidak ada data").border = thin
        ws2.cell(2, 2, 0).border = thin
    ws2.column_dimensions["A"].width = 36
    ws2.column_dimensions["B"].width = 12

    # Sheet 3: Volume mingguan
    ws3 = wb.create_sheet("Volume Mingguan")
    ws3.append(["Minggu", "Dibuat", "Diselesaikan"])
    for col in (1, 2, 3):
        cell = ws3.cell(1, col)
        cell.fill = header_fill
        cell.font = header_font
        cell.border = thin
    for i, week in enumerate(data.get("weeks") or [], start=2):
        ws3.cell(i, 1, week["label"]).border = thin
        ws3.cell(i, 2, week["created"]).border = thin
        ws3.cell(i, 3, week["resolved"]).border = thin
    ws3.column_dimensions["A"].width = 22
    ws3.column_dimensions["B"].width = 12
    ws3.column_dimensions["C"].width = 14

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def search_services(db: Session, q: str | None) -> list[Service]:
    query = db.query(Service).filter(Service.is_active.is_(True))
    if q:
        query = query.filter(Service.name.ilike(f"%{q.strip()}%"))
    return query.order_by(Service.name.asc()).all()


def search_categories(db: Session, service_id: int, q: str | None) -> list[Category]:
    query = db.query(Category).filter(Category.service_id == service_id)
    if q:
        query = query.filter(Category.name.ilike(f"%{q.strip()}%"))
    return query.order_by(Category.name.asc()).all()


def service_item(service: Service) -> dict:
    categories = sorted(service.categories or [], key=lambda row: (row.name or "").lower())
    return {
        "id": service.id,
        "name": service.name,
        "is_active": bool(service.is_active),
        "categories": [
            {
                "id": row.id,
                "name": row.name,
                "color": row.color,
                "service_id": row.service_id,
            }
            for row in categories
        ],
    }


def list_services_manage(db: Session, q: str | None = None) -> list[Service]:
    query = db.query(Service).options(joinedload(Service.categories))
    if q:
        query = query.filter(Service.name.ilike(f"%{q.strip()}%"))
    return query.order_by(Service.name.asc()).all()


def get_service_or_404(db: Session, service_id: int) -> Service:
    row = (
        db.query(Service)
        .options(joinedload(Service.categories))
        .filter(Service.id == service_id)
        .first()
    )
    if not row:
        raise NotFoundException("Layanan tidak ditemukan")
    return row


def create_service(db: Session, *, name: str, is_active: bool = True) -> Service:
    clean = (name or "").strip()
    if not clean:
        raise BadRequestException("Nama layanan wajib diisi")
    exists = db.query(Service).filter(Service.name == clean).first()
    if exists:
        raise BadRequestException("Nama layanan sudah digunakan")
    row = Service(name=clean, is_active=is_active)
    db.add(row)
    db.commit()
    db.refresh(row)
    return get_service_or_404(db, row.id)


def update_service(
    db: Session,
    service_id: int,
    *,
    name: str,
    is_active: bool,
) -> Service:
    row = get_service_or_404(db, service_id)
    clean = (name or "").strip()
    if not clean:
        raise BadRequestException("Nama layanan wajib diisi")
    exists = (
        db.query(Service)
        .filter(Service.name == clean, Service.id != service_id)
        .first()
    )
    if exists:
        raise BadRequestException("Nama layanan sudah digunakan")
    row.name = clean
    row.is_active = is_active
    db.commit()
    return get_service_or_404(db, service_id)


def delete_service(db: Session, service_id: int) -> None:
    row = get_service_or_404(db, service_id)
    in_use = db.query(Ticket.id).filter(Ticket.service_id == service_id).first()
    if in_use:
        row.is_active = False
        db.commit()
        return
    db.delete(row)
    db.commit()


def create_category(
    db: Session,
    service_id: int,
    *,
    name: str,
    color: str = "#8b1538",
) -> Category:
    get_service_or_404(db, service_id)
    clean = (name or "").strip()
    if not clean:
        raise BadRequestException("Nama kategori wajib diisi")
    exists = (
        db.query(Category)
        .filter(Category.service_id == service_id, Category.name == clean)
        .first()
    )
    if exists:
        raise BadRequestException("Kategori sudah ada pada layanan ini")
    row = Category(service_id=service_id, name=clean, color=(color or "#8b1538").strip())
    db.add(row)
    db.commit()
    db.refresh(row)
    return row


def update_category(
    db: Session,
    category_id: int,
    *,
    name: str,
    color: str,
) -> Category:
    row = db.query(Category).filter(Category.id == category_id).first()
    if not row:
        raise NotFoundException("Kategori tidak ditemukan")
    clean = (name or "").strip()
    if not clean:
        raise BadRequestException("Nama kategori wajib diisi")
    exists = (
        db.query(Category)
        .filter(
            Category.service_id == row.service_id,
            Category.name == clean,
            Category.id != category_id,
        )
        .first()
    )
    if exists:
        raise BadRequestException("Kategori sudah ada pada layanan ini")
    row.name = clean
    row.color = (color or row.color or "#8b1538").strip()
    db.commit()
    db.refresh(row)
    return row


def delete_category(db: Session, category_id: int) -> None:
    row = db.query(Category).filter(Category.id == category_id).first()
    if not row:
        raise NotFoundException("Kategori tidak ditemukan")
    in_use = db.query(Ticket.id).filter(Ticket.category_id == category_id).first()
    if in_use:
        raise BadRequestException("Kategori masih dipakai tiket dan tidak bisa dihapus")
    db.delete(row)
    db.commit()
